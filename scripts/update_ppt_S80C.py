"""
Insert one slide for S80C (Cell wall genes with consistent GRBV response,
focused subset of S80B) immediately after the existing S80B slide.

Run: python3 scripts/update_ppt_S80C.py
"""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = Path("/Users/prempratapsingh/Data_Analysis")
PPT  = ROOT / "GRBV_Funding_Committee_Presentation.pptx"
FIGS = ROOT / "figures" / "supp"

TEMPLATE_INDEX = 27
SECTION_LABEL  = "PLUS: agenda gap"
SECTION_ROOT   = "6 / Gap-fill"
ANCHOR_TITLE   = "PLUS (per-gene, S80B)"

NAVY  = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY  = RGBColor(0x64, 0x74, 0x8B)

NEW_SLIDE = {
    "title": "PLUS (per-gene, S80C): Cell-wall genes with consistent GRBV response (focused subset)",
    "image": "S80C_cell_wall_focused.png",
    "body": [
        "Companion to S80B (full 161-gene landscape).  S80C filters down to 21 cell-wall genes that show a substantive, consistent GRBV response across season x stage cells -- a focused panel readable at full resolution with cell-text annotation.",
        "Selection (genes retained if EITHER): (i) |mean log2FC| >= 1 in >=4 of 8 (season x stage) cells (the strict S80B labelled set; 13 genes), OR (ii) |mean log2FC| >= 0.45 in >=6 of 8 cells (broader consistent-response set; 8 additional genes).  The broad amplitude was auto-tuned because cell-wall genes respond more mutedly than the phenylpropanoid + PR set.",
        "Annotation: row labels are 'Symbol — short_function' from NCBI gene-info where available, 'LOC###### — best NCBI description' otherwise.  No bare LOC IDs.",
        "Layout: same season x stage facet structure as S80B, but with sub-categories rendered via ggh4x::scale_y_discrete(guide = 'axis_nested') as a left-side nested y-axis label (Pectin -> Hemicellulose -> Cellulose -> Lignin -> Expansin).  Within each sub-category, rows are clustered (Euclidean / Ward.D2).",
        "Cell-text overlay: |log2FC| >= 1.5 cells display the value to one decimal place (white text when |log2FC| > 2, black otherwise).  Color scale capped at +/- 3 with scales::squish.  10 x 9 in @ 600 DPI.",
    ],
}


def set_text(tf, text, *, size=12, bold=False, color=None, name="Cambria"):
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = name
    if color:
        r.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor(*color)


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def find_slide_by_title(prs, title_fragment):
    for i, s in enumerate(prs.slides):
        t = find_shape(s, "Text 1")
        if t is not None and title_fragment.lower() in t.text_frame.text.lower():
            return i
    return None


def clone_slide(prs, template_idx):
    src = prs.slides[template_idx]
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    skip = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
    for child in src.shapes._spTree:
        if child.tag in skip:
            continue
        new.shapes._spTree.append(deepcopy(child))
    return new


def replace_picture(slide, new_image_path):
    for sh in slide.shapes:
        if sh.shape_type == 13:
            left, top, w, h = sh.left, sh.top, sh.width, sh.height
            sp = sh._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(str(new_image_path), left, top, width=w, height=h)
            return
    raise RuntimeError("no picture shape found in template")


def replace_body_text(slide, bullets):
    body = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Text") and sh.name not in (
                "Text 1", "Text 3", "Text 6", "Text 7"):
            if body is None or sh.width * sh.height > body.width * body.height:
                body = sh
    if body is None:
        raise RuntimeError("no body text frame found")
    tf = body.text_frame
    tf.clear(); tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run(); r.text = line
        r.font.size = Pt(10); r.font.name = "Cambria"
        r.font.color.rgb = NAVY


def update_title(slide, title):
    t = find_shape(slide, "Text 1")
    if t is not None:
        set_text(t.text_frame, title, size=16, bold=True, color=WHITE)


def update_section_label(slide, label, root_label):
    lbl = find_shape(slide, "Text 3")
    if lbl is not None:
        set_text(lbl.text_frame, label, size=9, bold=True, color=WHITE)
    root = find_shape(slide, "Text 6")
    if root is not None:
        set_text(root.text_frame, root_label, size=8)


def reorder_slide(prs, from_idx, to_idx):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moving = slides[from_idx]
    xml_slides.remove(moving)
    xml_slides.insert(to_idx, moving)


def main():
    prs = Presentation(str(PPT))
    total_before = len(prs.slides)
    print(f"Opened presentation with {total_before} slides.")

    anchor_idx = find_slide_by_title(prs, ANCHOR_TITLE)
    if anchor_idx is None:
        raise SystemExit(f"Could not find anchor slide '{ANCHOR_TITLE}'")
    print(f"  inserting after slide {anchor_idx + 1} ('{ANCHOR_TITLE}')")

    new = clone_slide(prs, TEMPLATE_INDEX)
    update_title(new, NEW_SLIDE["title"])
    update_section_label(new, SECTION_LABEL, SECTION_ROOT)
    img = FIGS / NEW_SLIDE["image"]
    if not img.exists():
        raise FileNotFoundError(img)
    replace_picture(new, img)
    replace_body_text(new, NEW_SLIDE["body"])
    appended_idx = len(prs.slides) - 1
    reorder_slide(prs, appended_idx, anchor_idx + 1)
    print(f"  cloned + moved to position {anchor_idx + 2}")

    # Refresh footers across the whole deck
    total_after = len(prs.slides)
    for i, s in enumerate(prs.slides, start=1):
        for shape_name in ("Text 5", "Text 7"):
            num = find_shape(s, shape_name)
            if num is not None and "/" in num.text_frame.text:
                cur_size = None
                for p in num.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cur_size = r.font.size.pt; break
                    if cur_size: break
                cur_size = cur_size or 8
                set_text(num.text_frame, f"{i} / {total_after}",
                         size=cur_size, name="Cambria",
                         color=GREY if shape_name == "Text 5" else None)

    # Update slide-4 roadmap if section ranges shifted
    slide4 = prs.slides[3]
    s7_idx = find_slide_by_title(prs, "Integrated disease cascade")
    s8_idx = find_slide_by_title(prs, "Reanalysis, deliverables")
    if s7_idx is not None and s8_idx is not None:
        s7_label = f"slides {s7_idx + 1} – {s8_idx}"
        s8_label = f"slides {s8_idx + 1} – {total_after}"
        for shape_name, label in (("Text 33", s7_label), ("Text 37", s8_label)):
            sh = find_shape(slide4, shape_name)
            if sh is not None:
                set_text(sh.text_frame, label, size=13, bold=False, name="Cambria", color=NAVY)
                print(f"  slide 4 {shape_name} -> {label!r}")

    prs.save(str(PPT))
    print(f"Saved {PPT} with {total_after} slides (+{total_after - total_before}).")


if __name__ == "__main__":
    main()
