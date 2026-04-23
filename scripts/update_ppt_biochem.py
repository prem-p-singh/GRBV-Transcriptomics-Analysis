"""
Update GRBV_Funding_Committee_Presentation.pptx with the biochemistry
integration: add three PLUS(biochem) slides between Section 6 and Section 7,
and reframe SPS1 references on the existing mediation and synthesis slides.

Run: python3 scripts/update_ppt_biochem.py
"""
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path('/Users/prempratapsingh/Data_Analysis')
PPT  = ROOT / 'GRBV_Funding_Committee_Presentation.pptx'
FIGS = ROOT / 'figures' / 'supp'

TEMPLATE_INDEX = 27     # slide 28 (1-based): "Direct vs indirect temp effects" — nice 2-col layout
SECTION_LABEL  = 'PLUS: biochemistry'
SECTION_ROOT   = '4 / Results'

BIOCHEM_SLIDES = [
    {
        'title': 'PLUS (biochemistry): juice chemistry mirrors the inverted trajectory',
        'image': 'S63_biochem_season_infection.png',
        'body': [
            'Harvest juice-chemistry panel on n = 33 paired samples (2021 = 15, 2022 = 18): Brix, pH, TA, malic, NOPA, NH3, YAN. Two-way Season x Infection ANOVA.',
            'Titratable acidity is the cleanest phenotypic mirror of the transcriptomic inverted seasonal trajectory: no main effect of season (p = 0.69) or infection (p = 0.16), but Season x Infection interaction eta^2 = 0.77 (p = 4.3e-11).',
            'GRBV effect on TA: +0.62 g/L in 2021 cool, -0.79 g/L in 2022 hot. Malic acid follows the same inversion (+398 mg/L vs -195 mg/L).',
            'Brix suppression is significant in both seasons but 2.6 times stronger in 2022 (-9.83) than 2021 (-3.77). Interaction eta^2 = 0.13.',
            'pH: season-dominated (eta^2 = 0.89); no Season x Infection interaction, so the infected-berry pH drop is constant across years.',
        ],
    },
    {
        'title': 'Cohort trajectory at harvest: biphasic recovery in 2021',
        'image': 'S62_biochem_cohort_trajectory.png',
        'body': [
            'RB(+) samples regressed against years since infection per season, dashed line = RB(-) mean.',
            '2021 Brix rises +0.10 degBx per year with YSI (r = 0.76, p = 0.0045). Older infections partially recover sugar.',
            'In parallel: 2021 NOPA declines -2.45 mg/L / yr (p = 0.006), NH3 -2.40 mg/L / yr (p = 0.035), YAN -4.38 mg/L / yr (p = 0.07).',
            'Phenotypic reading: cool-season vines compensate on sugar over infection age at the cost of durable nitrogen-metabolism depletion.',
            'This is the direct biochemical analogue of the Q1 biphasic transcriptomic recovery. 2022 shows same directions without statistical support, consistent with hot-season dynamics overriding cohort-age effects.',
        ],
    },
    {
        'title': 'SPS1 reframed: compensatory, not rate-limiting',
        'image': 'S61_biochem_validation.png',
        'body': [
            'Per-sample expression (log2 CPM) against harvest juice chemistry; partial Spearman rho controls for season.',
            'SPS1 (LOC100232974) correlates negatively with Brix within-season (2021 rho = -0.73, p = 0.002; partial rho = -0.42). Higher SPS1 in low-Brix infected berries.',
            'Interpret as compensatory upregulation of sucrose synthesis in metabolically-stressed berries, not the rate-limiting step. Phloem loading or transport is the plausible failing node.',
            'LOC100266309 is the strongest within-season host validation (partial rho = -0.45, 2021 rho = -0.65). MBF1c has no juice-chemistry correlate at harvest.',
            'Viral titer x Brix is a 2022-specific association (partial rho = -0.25, 2022 rho = -0.44) rather than a universal relationship.',
            'The mediation status of SPS1 (HIMA + main-paper sweep) is unchanged; only the "direct molecular link to sugar phenotype" interpretation is retracted.',
        ],
    },
]


def set_text(text_frame, text, size=12, bold=False, color=None, bullet=False):
    """Overwrite a text frame with a single paragraph."""
    text_frame.clear()
    p = text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def clone_slide(prs, template_idx):
    """Clone a slide by deep-copying its XML spTree and re-linking picture rels."""
    src = prs.slides[template_idx]
    new = prs.slides.add_slide(src.slide_layout)

    # Remove default placeholders on the fresh slide
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)

    # Deep-copy every child of src spTree
    src_spTree = src.shapes._spTree
    dst_spTree = new.shapes._spTree
    skip_tags = {qn('p:nvGrpSpPr'), qn('p:grpSpPr')}
    for child in src_spTree:
        if child.tag in skip_tags:
            continue
        dst_spTree.append(deepcopy(child))

    # Fix picture rId references: python-pptx assigns new relationship ids,
    # so any cloned <p:blipFill><a:blip r:embed="rIdN"/> will be dangling.
    # Walk copied pics, import the underlying image from the src rel, and
    # rewrite the embed rId to the new one.
    for new_pic in dst_spTree.findall(qn('p:pic')):
        blip = new_pic.find('.//' + qn('a:blip'))
        if blip is None:
            continue
        old_rId = blip.get(qn('r:embed'))
        if not old_rId:
            continue
        try:
            src_image = src.part.related_part(old_rId).blob
            image_part, new_rId = new.part.get_or_add_image(
                image_part_or_stream(src_image))
            blip.set(qn('r:embed'), new_rId)
        except Exception as e:
            print(f'    warn: could not relink picture rId={old_rId}: {e}')

    return new


def image_part_or_stream(blob):
    from io import BytesIO
    return BytesIO(blob)


def replace_picture(slide, new_image_path):
    """Swap the first picture on `slide` for a new image, preserving box."""
    for sh in slide.shapes:
        if sh.shape_type == 13:       # PICTURE
            left, top, width, height = sh.left, sh.top, sh.width, sh.height
            sp = sh._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(str(new_image_path), left, top,
                                     width=width, height=height)
            return
    raise RuntimeError('no picture shape found')


def replace_body_text(slide, bullets):
    """Rewrite the right-column body text block with new bullets."""
    # Body text shape is the widest text block not named 'Text 1/3/6/7'.
    body = None
    for sh in slide.shapes:
        if sh.name in ('Text 9', 'Text 10', 'Text 11', 'Text 12'):
            body = sh
            break
    if body is None:
        # fall back: largest text frame after title
        candidates = [sh for sh in slide.shapes if sh.has_text_frame
                      and sh.name not in ('Text 1', 'Text 3', 'Text 6', 'Text 7')]
        candidates.sort(key=lambda s: s.width * s.height, reverse=True)
        body = candidates[0] if candidates else None
    if body is None:
        raise RuntimeError('no body text frame found')

    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(10)
        run.font.name = 'Arial'


def update_slide_number(slide, n, total):
    num = find_shape(slide, 'Text 7')
    if num is not None:
        set_text(num.text_frame, f'{n} / {total}', size=8)


def update_section_label(slide, label, root_label):
    lbl = find_shape(slide, 'Text 3')
    if lbl is not None:
        set_text(lbl.text_frame, label, size=9, bold=True, color=(255, 255, 255))
    root = find_shape(slide, 'Text 6')
    if root is not None:
        set_text(root.text_frame, root_label, size=8)


def update_title(slide, title):
    t = find_shape(slide, 'Text 1')
    if t is not None:
        set_text(t.text_frame, title, size=16, bold=True)


def reorder_slide(prs, from_idx, to_idx):
    """Move slide at from_idx to to_idx (both zero-based)."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moving = slides[from_idx]
    xml_slides.remove(moving)
    xml_slides.insert(to_idx, moving)


def reframe_existing_slides(prs):
    """Update SPS1 references on slides 26 and 50 per the biochem reframe."""
    # Slide 26 (0-indexed 25) - mediation intro
    s26 = prs.slides[25]
    body_shape = find_shape(s26, 'Text 9') or None
    if body_shape is None:
        for sh in s26.shapes:
            if sh.has_text_frame and 'Causal-mediation sweep' in sh.text_frame.text:
                body_shape = sh
                break
    if body_shape is not None:
        new_bullets = [
            'Causal-mediation sweep on the top 2,000 variable genes: 39 joint-significant Temp -> Titer -> Gene mediators at BH FDR < 0.05 on both paths.',
            'Top hit: LOC100249249, annotated by NCBI as multiprotein-bridging factor 1c (MBF1c), a known plant HSF co-activator. Joint p = 4.8e-11; titer mediates 61.8 % of the temperature effect on MBF1c.',
            'Second-tier mediator: LOC100232974 = sucrose-phosphate synthase 1 (SPS1). Paired harvest biochemistry (S61) shows SPS1 correlates NEGATIVELY with Brix within-season (2021 rho = -0.73), so SPS1 is upregulated in low-Brix infected berries. Compensatory, not rate-limiting.',
            'Third robust mediator: LOC100266309 (uncharacterised). Strongest within-season host biochemistry validation: partial rho vs Brix = -0.45.',
            'Category breakdown: 39 titer-mediated, 520 temperature-only (direct effect without titer as conduit), 1,441 non-significant. At the old tv = 500 threshold, zero genes passed joint significance.',
        ]
        tf = body_shape.text_frame
        tf.clear(); tf.word_wrap = True
        for i, line in enumerate(new_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            r = p.add_run()
            r.text = line
            r.font.size = Pt(10); r.font.name = 'Arial'
        print('  reframed slide 26 mediation bullets')

    # Slide 50 (0-indexed 49) - synthesis cascade Station 4
    s50 = prs.slides[49]
    for sh in s50.shapes:
        if sh.has_text_frame and 'Station 4' in sh.text_frame.text:
            body = sh
            break
    else:
        body = None
    if body is not None:
        tf = body.text_frame
        paras = list(tf.paragraphs)
        for p in paras:
            txt = p.text
            if 'Station 4' in txt and 'SPS1' in txt:
                # Rewrite this paragraph
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
                r = p.add_run()
                r.text = ('Station 4, Sugar metabolism rewired: DOF3.6 and DOF1.7 motifs '
                          'enriched in brown + yellow modules; Brix suppressed 4 to 9.8 degBx in infected '
                          'berries, with TA and malic inverted between 2021 cool and 2022 hot seasons (S63); '
                          'top biomarker maps to sugar QTL on chr 1. SPS1 is a robust temperature mediator '
                          'but is upregulated in low-Brix infected berries (S61), so its role is compensatory '
                          'rather than rate-limiting.')
                r.font.size = Pt(10); r.font.name = 'Arial'
                print('  reframed slide 50 Station 4 text')
                break


def add_biochem_correction_bullet(prs):
    """On slide 52 (reanalysis corrected), append a new bullet about biochem + SPS1 reframe."""
    s52 = prs.slides[51]
    body = None
    for sh in s52.shapes:
        if sh.has_text_frame and 'Universal DEGs revised' in sh.text_frame.text:
            body = sh; break
    if body is None:
        print('  could not find slide 52 body; skipping')
        return
    tf = body.text_frame
    p = tf.add_paragraph()
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = ('SPS1 role reframed: paired harvest biochemistry (n = 33, Figs S61-S63) shows SPS1 is '
              'upregulated in low-Brix infected berries (2021 rho = -0.73). Mediator status stands; '
              'the "direct molecular link to sugar phenotype" interpretation is retracted. Season x '
              'Infection interaction on TA (eta^2 = 0.77) is the new phenotypic mirror of the inverted '
              'seasonal trajectory.')
    r.font.size = Pt(10); r.font.name = 'Arial'
    print('  added biochem correction to slide 52')


def main():
    prs = Presentation(str(PPT))
    total_before = len(prs.slides)
    print(f'Opened presentation with {total_before} slides.')

    # Reframe existing slides (do this before adding new ones so indices are stable)
    reframe_existing_slides(prs)
    add_biochem_correction_bullet(prs)

    # Clone template slide 3 times, customise, then reorder to sit before slide 49 (SECTION 7)
    added_indices = []
    for entry in BIOCHEM_SLIDES:
        new = clone_slide(prs, TEMPLATE_INDEX)
        update_title(new, entry['title'])
        update_section_label(new, SECTION_LABEL, SECTION_ROOT)
        img_path = FIGS / entry['image']
        if not img_path.exists():
            raise FileNotFoundError(img_path)
        replace_picture(new, img_path)
        replace_body_text(new, entry['body'])
        added_indices.append(len(prs.slides) - 1)
        print(f"  cloned template -> new slide '{entry['title'][:60]}...'")

    # Reorder: move the three newly-appended slides to positions 48, 49, 50
    # (one after the last viral-biology slide, before the SECTION 7 header).
    target_positions = [48, 49, 50]
    for new_idx, target in zip(added_indices, target_positions):
        current_idx = len(prs.slides) - (len(added_indices) - added_indices.index(new_idx))
        reorder_slide(prs, current_idx, target)

    # Refresh slide-number footers across the whole deck
    total_after = len(prs.slides)
    for i, s in enumerate(prs.slides, start=1):
        num = find_shape(s, 'Text 7')
        if num is not None and '/' in num.text_frame.text:
            set_text(num.text_frame, f'{i} / {total_after}', size=8)

    prs.save(str(PPT))
    print(f'Saved {PPT} with {total_after} slides (+{total_after - total_before}).')


if __name__ == '__main__':
    main()
