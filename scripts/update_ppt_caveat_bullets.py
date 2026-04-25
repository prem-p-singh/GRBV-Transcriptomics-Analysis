"""
Append two caveat bullets surfaced by the agenda-gap pass:
  - Slide 30 (Q4 mechanism MBF1c): note that S79's Bayesian random-intercept
    posterior softens MBF1c.
  - Slide 63 (New dimension 2 — three temporal phases): note that S78's
    segmented regression rejects discrete phase transitions; framing is
    descriptive only.

Run: python3 scripts/update_ppt_caveat_bullets.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

ROOT = Path("/Users/prempratapsingh/Data_Analysis")
PPT  = ROOT / "GRBV_Funding_Committee_Presentation.pptx"
NAVY = RGBColor(0x0F, 0x17, 0x2A)

# (1-indexed slide number, exact title fragment used to anchor by content,
#  caveat bullet)
EDITS = [
    {
        "slide_idx": 30,
        "title_anchor": "Q4 mechanism",
        "bullet": (
            "Caveat (added 2026-04-24, see S79): a Bayesian random-intercept reanalysis "
            "(brms y ~ Inf + Season + Stage + (1 | vine_id), 4 chains x 2000 iter) "
            "softens MBF1c to posterior mean -0.25 [-0.81, +0.33] with 19% in ROPE. "
            "The other 5 panel genes hold up under the more conservative inference "
            "(RdRp1 +1.78, F-box +2.01, glycine-rich -1.33, LOC100266309 +0.80, "
            "SPS1 +0.52, all with 95% CI outside ROPE).  The headline mediator's "
            "effect direction needs qPCR / VIGS to confirm before submission."
        ),
    },
    {
        "slide_idx": 63,
        "title_anchor": "New dimension 2",
        "bullet": (
            "Caveat (added 2026-04-24, see S78): segmented regression on years-since-"
            "infection rejects discrete phase transitions.  best_k = 0 (linear) for "
            "ALL 15 responses (PC1, PC2, 13 module eigengenes) by BIC.  Treat acute / "
            "peak / chronic as descriptive cohort-bin labels, not statistically "
            "supported phase architecture.  The continuous-gradient model is "
            "sufficient and simpler; no mechanistic phase-transition needed."
        ),
    },
]


def append_bullet(slide, text):
    body = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Text") and sh.name not in (
                "Text 1", "Text 3", "Text 4", "Text 5", "Text 6", "Text 7"):
            if body is None or sh.width * sh.height > body.width * body.height:
                body = sh
    if body is None:
        raise RuntimeError("no body text frame found")
    tf = body.text_frame
    p = tf.add_paragraph(); p.space_after = Pt(4)
    r = p.add_run(); r.text = text
    r.font.size = Pt(10); r.font.name = "Cambria"; r.font.color.rgb = NAVY


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def main():
    prs = Presentation(str(PPT))
    print(f"Opened presentation with {len(prs.slides)} slides.")
    for edit in EDITS:
        sn = edit["slide_idx"]
        slide = prs.slides[sn - 1]
        title_sh = find_shape(slide, "Text 1")
        if title_sh is None or edit["title_anchor"].lower() not in title_sh.text_frame.text.lower():
            print(f"  warn: slide {sn} title does not match anchor {edit['title_anchor']!r}")
            continue
        append_bullet(slide, edit["bullet"])
        print(f"  appended caveat bullet to slide {sn}")
    prs.save(str(PPT))
    print("Saved.")


if __name__ == "__main__":
    main()
