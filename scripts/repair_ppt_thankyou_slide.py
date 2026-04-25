"""
Repair the duplicate-pointer corruption introduced during the
update_ppt_split_S80.py reorder operation: 'rId72' in
presentation.xml.rels was pointing at the same file (slide82.xml) that
'rId87' already references, causing slide 82 in the deck to be a
duplicate of slide 67 (S80A) and the original 'Thank you' slide content
to be lost from the active file system.

Recovery: extract the original Thank You slide.xml from the
pre-S80-split backup (/tmp/GRBV_pre_S80split_201143.pptx, where it
lived at slide78.xml referenced by rId72 at sldIdLst position 80), inject
it as ppt/slides/slide83.xml in the current file, add the corresponding
rels (with the broken notesSlide55.xml reference stripped), add a
Content_Types Override, and rewire rId72 to point at the new
slide83.xml.

A backup is taken first.

Run: python3 scripts/repair_ppt_thankyou_slide.py
"""
from pathlib import Path
import zipfile
import shutil
import sys
import re
from datetime import datetime

PPT    = Path("/Users/prempratapsingh/Data_Analysis/GRBV_Funding_Committee_Presentation.pptx")
BACKUP = Path("/tmp/GRBV_pre_S80split_201143.pptx")
NEW_SLIDE_NAME = "slide83.xml"   # Pick the next free number; current files go up to slide82.xml.


def diagnose_current(p):
    with zipfile.ZipFile(p) as z:
        pres = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
    sld_rids = re.findall(r'<p:sldId\s+id="\d+"\s+r:id="(rId\d+)"', pres)
    rid_target = dict(re.findall(
        r'Id="(rId\d+)"\s+Type="[^"]+slide"\s+Target="([^"]+)"', rels))
    pos_files = [rid_target.get(rid) for rid in sld_rids]
    from collections import Counter
    dup = {f: c for f, c in Counter(pos_files).items() if c > 1}
    if not dup:
        sys.exit("abort: no duplicate slide-file pointers found; nothing to repair")
    print(f"  duplicates: {dup}")
    # The orphan-pointer slide is at the position whose rId is the SECOND occurrence
    pos_rids = []
    for fname in dup:
        positions_rids = [(i+1, sld_rids[i]) for i, v in enumerate(pos_files) if v == fname]
        # First occurrence is the "real" inserted slide; second is the orphan we need to fix.
        first_pos, first_rid = positions_rids[0]
        for pos, rid in positions_rids[1:]:
            print(f"  orphan duplicate: pos {pos} (rId {rid}) -> {fname} "
                  f"(real owner: pos {first_pos} via {first_rid})")
            pos_rids.append((pos, rid, fname))
    return pos_rids


def extract_thankyou_from_backup():
    with zipfile.ZipFile(BACKUP) as z:
        pres = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
    sld_rids = re.findall(r'<p:sldId\s+id="\d+"\s+r:id="(rId\d+)"', pres)
    rid_target = dict(re.findall(
        r'Id="(rId\d+)"\s+Type="[^"]+slide"\s+Target="([^"]+)"', rels))
    # Find the slide with "Thank you" content
    with zipfile.ZipFile(BACKUP) as z:
        for rid in sld_rids:
            target = rid_target[rid]
            slide_path = "ppt/" + target.replace("../", "")
            xml = z.read(slide_path).decode()
            if "Thank you" in xml and "Questions and discussion" in xml:
                rels_path = (slide_path
                             .replace("ppt/slides/", "ppt/slides/_rels/")
                             + ".rels")
                rels_xml = z.read(rels_path).decode()
                print(f"  found Thank You slide content at {slide_path}")
                return xml, rels_xml
    sys.exit("abort: 'Thank you' slide not found in backup")


def strip_notes_ref(rels_xml):
    """Remove any <Relationship .../> that references a notesSlide; the
    referenced notes file may not exist in the current pptx."""
    return re.sub(
        r'<Relationship\s+Id="rId\d+"\s+Type="[^"]+notesSlide"\s+Target="[^"]+"\s*/>',
        "", rels_xml)


def repair(p, orphan_pos, orphan_rid, slide_xml, rels_xml):
    backup = p.with_suffix(
        f".pre_thankyou_repair_{datetime.now():%Y%m%d_%H%M%S}.pptx")
    shutil.copy2(p, backup)
    print(f"  backup -> {backup}")

    rels_xml_clean = strip_notes_ref(rels_xml)
    new_slide_full = f"ppt/slides/{NEW_SLIDE_NAME}"
    new_rels_full  = f"ppt/slides/_rels/{NEW_SLIDE_NAME}.rels"

    tmp_out = p.with_suffix(".pptx.tmp")
    rels_modified = ct_modified = False

    with zipfile.ZipFile(p, "r") as zin, \
         zipfile.ZipFile(tmp_out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename == "ppt/_rels/presentation.xml.rels":
                txt = data.decode()
                # Repoint orphan_rid from slides/slide82.xml -> slides/<NEW>
                pattern = (rf'(<Relationship\s+Id="{re.escape(orphan_rid)}"'
                           rf'\s+Type="[^"]+slide"\s+Target=")slides/slide82\.xml"')
                new_txt, n = re.subn(pattern,
                                     rf'\1slides/{NEW_SLIDE_NAME}"',
                                     txt)
                if n != 1:
                    sys.exit(f"abort: could not repoint {orphan_rid} (matched {n} times)")
                data = new_txt.encode()
                rels_modified = True

            elif item.filename == "[Content_Types].xml":
                txt = data.decode()
                # Add Override for the new slide before </Types>
                override = (f'<Override PartName="/ppt/slides/{NEW_SLIDE_NAME}"'
                            f' ContentType="application/vnd.openxmlformats-'
                            f'officedocument.presentationml.slide+xml"/>')
                if override in txt:
                    print("  CT override already present; skipping")
                else:
                    new_txt = txt.replace("</Types>", override + "</Types>")
                    if new_txt == txt:
                        sys.exit("abort: failed to inject CT override")
                    data = new_txt.encode()
                ct_modified = True

            zout.writestr(item, data)

        # Append the new slide files
        zout.writestr(new_slide_full, slide_xml)
        zout.writestr(new_rels_full,  rels_xml_clean)

    if not (rels_modified and ct_modified):
        sys.exit("abort: incomplete repair")

    tmp_out.replace(p)
    print(f"  injected {new_slide_full} ({len(slide_xml)} bytes) + rels")


def verify(p):
    with zipfile.ZipFile(p) as z:
        pres = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
        ct   = z.read("[Content_Types].xml").decode()
    sld_rids = re.findall(r'<p:sldId\s+id="\d+"\s+r:id="(rId\d+)"', pres)
    rid_target = dict(re.findall(
        r'Id="(rId\d+)"\s+Type="[^"]+slide"\s+Target="([^"]+)"', rels))
    pos_files = [rid_target.get(rid) for rid in sld_rids]
    from collections import Counter
    dup = {f: c for f, c in Counter(pos_files).items() if c > 1}
    assert not dup, f"duplicates remain: {dup}"
    assert f"PartName=\"/ppt/slides/{NEW_SLIDE_NAME}\"" in ct, \
        "CT override missing"
    # Open with python-pptx to confirm structural validity
    from pptx import Presentation
    pp = Presentation(str(p))
    n = len(pp.slides)
    last = pp.slides[n - 1]
    title = ""
    for sh in last.shapes:
        if sh.name == "Text 1" and sh.has_text_frame:
            title = sh.text_frame.text.strip(); break
    print(f"  verify: {n} slides, last slide title = {title!r}")
    assert "Thank you" in title, f"last slide is not 'Thank you' ({title!r})"


def main():
    print(f"Repairing {PPT}")
    orphan_pos_rids = diagnose_current(PPT)
    if len(orphan_pos_rids) != 1:
        sys.exit(f"abort: expected exactly 1 duplicate, got {len(orphan_pos_rids)}")
    orphan_pos, orphan_rid, _ = orphan_pos_rids[0]
    slide_xml, rels_xml = extract_thankyou_from_backup()
    repair(PPT, orphan_pos, orphan_rid, slide_xml, rels_xml)
    verify(PPT)
    print("Done.")


if __name__ == "__main__":
    main()
