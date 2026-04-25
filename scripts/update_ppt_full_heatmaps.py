"""
Add 2 dedicated full-resolution heatmap slides to GRBV_Funding_Committee_Presentation.pptx:

  S80 — Phenylpropanoid + PR + Cell-wall log2 FC, every season x cohort x stage
  S81 — Carbon translocation / phloem-loading module, same axes

Inserted after the rule 14/15 summary slides (currently slides 66 + 67) so the
full-resolution heatmap sits next to its summary view.

Run: python3 scripts/update_ppt_full_heatmaps.py
"""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = Path("/Users/prempratapsingh/Data_Analysis")
PPT  = ROOT / "GRBV_Funding_Committee_Presentation.pptx"
FIGS = ROOT / "figures" / "supp"

TEMPLATE_INDEX = 27
SECTION_LABEL  = "PLUS: agenda gap"
SECTION_ROOT   = "6 / Gap-fill"

NAVY  = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY  = RGBColor(0x64, 0x74, 0x8B)

NEW_SLIDES = [
    {
        "title": "PLUS (per-gene, full): Phenylpropanoid + PR + cell-wall log2 FC across season x cohort x stage",
        "image": "S80_pathway_full_heatmap.png",
        "anchor_after": "PLUS (per-gene path)",
        "body": [
            "Companion to S74's summary view.  All 252 curated genes plotted with full season x cohort x stage resolution.  Cell value = mean log2(CPM+1) RB(+) at (Sampling_Season, Maturity_stage, Years-Since-Infection cohort) - mean log2(CPM+1) RB(-) at matched (Sampling_Season, Maturity_stage).  RB(-) controls have no cohort, so the same baseline is reused across the 5 cohort columns within each (season, stage) facet.",
            "Three category panels stacked top-to-bottom: Phenylpropanoid (51), PR proteins (40), Cell wall (161).  Within each panel, genes are ordered by mean log2 FC (most-up at top).",
            "Layout: 8 outer facets (2 seasons x 4 maturity stages), each with 5 cohort columns inside.  40 condition cells per gene visible at one glance.",
            "Pattern check: most genes show stronger fold-changes at later stages (Post-veraison, Harvest) and in the 2022 hot season; this matches the inverted-trajectory + heat-amplification framing.",
            "Use as the in-paper Supplementary figure when reviewers ask 'show me the gene-by-gene picture' for the three agenda-named pathways.",
        ],
    },
    {
        "title": "PLUS (sugar transport, full): SWEET / SUT / STP / CWINV log2 FC across season x cohort x stage",
        "image": "S81_sugar_full_heatmap.png",
        "anchor_after": "PLUS (carbon translocation)",
        "body": [
            "Companion to S75's summary view.  All 33 expressed sugar-transport genes (SUT/SUC, STP/HXT, cell-wall + vacuolar invertases, sucrose synthases, SPS, hexokinase, trehalose, phloem markers) plotted with full season x cohort x stage resolution.  Same FC convention as S80.",
            "Subclass rows ordered functionally (source -> transport -> sink): SUT/SUC, STP/HXT, cell-wall INV, vacuolar INV, sucrose synthase, SPS, hexokinase, trehalose, phloem marker.",
            "Within-subclass gene order: by mean log2 FC across cells (most-up at top of each block).",
            "Reading: most subclasses show stage-driven patterns rather than infection-driven blocks; the panel-level result that the sugar module is 'stage-dominant, not viral' (rho_stage=0.94, no main season or cohort effect) is visible as the dominance of stage facets in the colour structure.",
            "Notable individual genes can still be picked out -- e.g., SPS1 (LOC100232974) and the strongest STP/HXT representatives -- and re-mapped to S61's Brix correlation panel for downstream interpretation.",
        ],
    },
]


def set_text(tf, text, *, size=12, bold=False, color=None, name="Cambria"):
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = name
    if color:
        r.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor(*color)


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def clone_slide(prs, template_idx):
    src = prs.slides[template_idx]
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    skip = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
    for child in src.shapes._spTree:
        if child.tag in skip:
            continue
        new.shapes._spTree.append(deepcopy(child))
    return new


def replace_picture(slide, new_image_path):
    for sh in slide.shapes:
        if sh.shape_type == 13:
            left, top, w, h = sh.left, sh.top, sh.width, sh.height
            sp = sh._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(str(new_image_path), left, top, width=w, height=h)
            return
    raise RuntimeError("no picture shape found in template")


def replace_body_text(slide, bullets):
    body = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Text") and sh.name not in (
                "Text 1", "Text 3", "Text 6", "Text 7"):
            if body is None or sh.width * sh.height > body.width * body.height:
                body = sh
    if body is None:
        raise RuntimeError("no body text frame found")
    tf = body.text_frame
    tf.clear(); tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run(); r.text = line
        r.font.size = Pt(10); r.font.name = "Cambria"
        r.font.color.rgb = NAVY


def update_title(slide, title):
    t = find_shape(slide, "Text 1")
    if t is not None:
        set_text(t.text_frame, title, size=16, bold=True, color=WHITE)


def update_section_label(slide, label, root_label):
    lbl = find_shape(slide, "Text 3")
    if lbl is not None:
        set_text(lbl.text_frame, label, size=9, bold=True, color=WHITE)
    root = find_shape(slide, "Text 6")
    if root is not None:
        set_text(root.text_frame, root_label, size=8)


def reorder_slide(prs, from_idx, to_idx):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    moving = slides[from_idx]
    xml_slides.remove(moving)
    xml_slides.insert(to_idx, moving)


def find_slide_by_title(prs, title_fragment):
    for i, s in enumerate(prs.slides):
        t = find_shape(s, "Text 1")
        if t is not None and title_fragment.lower() in t.text_frame.text.lower():
            return i
    return None


def main():
    prs = Presentation(str(PPT))
    total_before = len(prs.slides)
    print(f"Opened presentation with {total_before} slides.")

    # Insert each slide right after its anchor.  Insert in reverse order so
    # earlier insertions don't shift later anchor indices.
    insertions = []
    for entry in NEW_SLIDES:
        anchor_idx = find_slide_by_title(prs, entry["anchor_after"])
        if anchor_idx is None:
            print(f"  warn: anchor '{entry['anchor_after']}' not found; appending at end")
            anchor_idx = len(prs.slides) - 1
        insertions.append((anchor_idx, entry))

    # Sort by anchor descending so insertions don't shift indices for later items
    insertions.sort(key=lambda t: -t[0])

    for anchor_idx, entry in insertions:
        new = clone_slide(prs, TEMPLATE_INDEX)
        update_title(new, entry["title"])
        update_section_label(new, SECTION_LABEL, SECTION_ROOT)
        img = FIGS / entry["image"]
        if not img.exists():
            raise FileNotFoundError(img)
        replace_picture(new, img)
        replace_body_text(new, entry["body"])
        cur = len(prs.slides) - 1
        reorder_slide(prs, cur, anchor_idx + 1)
        print(f"  inserted '{entry['title'][:60]}...' after slide {anchor_idx + 1}")

    # Refresh footers across the whole deck
    total_after = len(prs.slides)
    for i, s in enumerate(prs.slides, start=1):
        for shape_name in ("Text 5", "Text 7"):
            num = find_shape(s, shape_name)
            if num is not None and "/" in num.text_frame.text:
                cur_size = None
                for p in num.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cur_size = r.font.size.pt; break
                    if cur_size: break
                cur_size = cur_size or 8
                set_text(num.text_frame, f"{i} / {total_after}",
                         size=cur_size, name="Cambria",
                         color=GREY if shape_name == "Text 5" else None)

    # Update slide-4 roadmap ranges (Section 7 + Section 8 shifted by +2)
    slide4 = prs.slides[3]
    # Determine current Section 7 and Section 8 start positions by title
    s7_idx = find_slide_by_title(prs, "Integrated disease cascade")
    s8_idx = find_slide_by_title(prs, "Reanalysis, deliverables")
    if s7_idx is not None and s8_idx is not None:
        s7_label = f"slides {s7_idx + 1} – {s8_idx}"
        s8_label = f"slides {s8_idx + 1} – {total_after}"
        for shape_name, label in (("Text 33", s7_label), ("Text 37", s8_label)):
            sh = find_shape(slide4, shape_name)
            if sh is not None:
                set_text(sh.text_frame, label, size=13, bold=False, name="Cambria", color=NAVY)
                print(f"  slide 4 {shape_name} -> {label!r}")

    prs.save(str(PPT))
    print(f"Saved {PPT} with {total_after} slides (+{total_after - total_before}).")


if __name__ == "__main__":
    main()
