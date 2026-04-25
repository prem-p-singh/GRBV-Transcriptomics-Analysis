"""
Refresh ONLY the Main 5 (mediation) + Main 10 (biology cascade) images on
the slides that already host them.  Caveat-bullet append for the Bayesian
softening was already done in scripts/update_ppt_caveat_bullets.py; this
script only swaps in the regenerated PNGs.

Run: python3 scripts/update_ppt_refresh_main510.py
"""
from pathlib import Path
from pptx import Presentation

ROOT = Path("/Users/prempratapsingh/Data_Analysis")
PPT  = ROOT / "GRBV_Funding_Committee_Presentation.pptx"

# Anchor by slide-title fragment, paired with the new image to embed
SWAPS = [
    {"title": "Q4 mechanism: Temp -> Titer -> Gene mediation",
     "img":   ROOT / "figures" / "main"  / "Main05_mediation.png"},
    {"title": "Synthesis: integrated molecular cascade",
     "img":   ROOT / "figures" / "main"  / "Main10_biology_cascade.png"},
]


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name: return sh
    return None


def main():
    prs = Presentation(str(PPT))
    print(f"Opened presentation with {len(prs.slides)} slides.")
    swaps_done = 0
    for entry in SWAPS:
        target = None
        for i, s in enumerate(prs.slides):
            t = find_shape(s, "Text 1")
            if t is not None and entry["title"].lower() in t.text_frame.text.lower():
                target = (i, s); break
        if target is None:
            print(f"  warn: anchor '{entry['title']}' not found"); continue
        idx, slide = target
        if not entry["img"].exists():
            print(f"  warn: missing image {entry['img']}"); continue
        pic = None
        for sh in slide.shapes:
            if sh.shape_type == 13:
                pic = sh; break
        if pic is None:
            print(f"  warn: no picture on slide {idx + 1}"); continue
        left, top, w, h = pic.left, pic.top, pic.width, pic.height
        pic._element.getparent().remove(pic._element)
        slide.shapes.add_picture(str(entry["img"]), left, top, width=w, height=h)
        print(f"  refreshed slide {idx + 1}: {entry['title'][:50]!r}")
        swaps_done += 1

    if swaps_done == 0:
        print("nothing to do"); return
    prs.save(str(PPT))
    print(f"Saved {PPT} ({swaps_done} image swap(s); slide count unchanged at {len(prs.slides)}).")


if __name__ == "__main__":
    main()
